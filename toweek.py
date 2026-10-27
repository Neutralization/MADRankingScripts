import json
import os
from math import floor

import arrow
import lxml.etree as etree  # noqa: PLR0402 - pyright cannot resolve "from lxml import etree"
import requests
from openpyxl import Workbook


def load_env_file():
    env_path = os.path.join(os.path.dirname(__file__), ".env")
    if not os.path.exists(env_path):
        return
    with open(env_path, encoding="utf-8") as env_file:
        for line in env_file:
            line = line.strip()
            if not line or line.startswith("#") or "=" not in line:
                continue
            key, value = line.split("=", 1)
            os.environ.setdefault(key.strip(), value.strip().strip('"\''))


load_env_file()
BASE_URL = os.environ["MAD_BILIRAN_BASE_URL"].rstrip("/")
COOKIE = os.environ["MAD_BILIRAN_COOKIE"]
DEFAULT_HEADERS = {
    "User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64; rv:109.0) Gecko/20100101 Firefox/115.0",
    "Accept": "text/html,application/xhtml+xml,application/xml;q=0.9,image/avif,image/webp,*/*;q=0.8",
    "Accept-Language": "zh-CN,zh;q=0.8,zh-TW;q=0.7,zh-HK;q=0.5,en-US;q=0.3,en;q=0.2",
    "DNT": "1",
    "Connection": "keep-alive",
    "Upgrade-Insecure-Requests": "1",
}
WEEKS = floor(
    (int(arrow.now("Asia/Shanghai").timestamp()) - 1428681600) / 3600 / 24 / 7
)


def getstatus(avid):
    with open("uids.json", encoding="utf-8-sig") as blacklist_file:
        blacklist = json.load(blacklist_file)
    response = requests.get(
        f"https://api.bilibili.com/x/web-interface/view?aid={avid}",
        headers=DEFAULT_HEADERS,
    )
    response.raise_for_status()
    result = response.json()
    if result["code"] == -404 or result["code"] == 62002 or result["code"] == 62012:
        return avid
    duration = result["data"]["duration"]
    if duration < 21:
        return avid
    mid = result["data"]["owner"]["mid"]
    if mid in blacklist:
        return avid
    return False


def create_session():
    """创建管理站点会话，供更新和下载操作共用。"""
    cookies = {"_biliran_sess": COOKIE}

    session = requests.Session()
    session.cookies.update(cookies)
    session.headers.update(DEFAULT_HEADERS)
    return session


def main():
    session = create_session()
    response = session.get(f"{BASE_URL}/mad_result_tmps")
    response.raise_for_status()

    html = etree.HTML(response.content, parser=etree.HTMLParser(encoding="utf-8"))
    if html is None:
        raise RuntimeError("列表页无法解析为 HTML")
    tokens = html.xpath('//meta[@name="csrf-token"]/@content')
    if not tokens or not tokens[0].strip():
        raise RuntimeError("列表页未提供 CSRF token，请检查登录会话是否有效")
    authenticity_token = tokens[0]
    ranking_rows = html.xpath(
        "//div[contains(@class, 'tile-body')]//table/tbody/tr[not(@class='danger')]"
    )

    avids = [int("".join(row.xpath("./td[3]/text()"))) for row in ranking_rows]

    excluded_avids = []
    for avid in avids:
        excluded_avid = getstatus(avid)
        if excluded_avid:
            excluded_avids.append(excluded_avid)

    historical_data = []
    for week in range(376, WEEKS):
        with open(f"./DATA/{week:03d}期数据.json", encoding="utf-8") as data_file:
            historical_data.extend(json.load(data_file))
    historical_avids = [
        item["av"] for item in historical_data if item["rank"] <= 20
    ]

    checklist = [avid for avid in avids if avid not in excluded_avids]

    print(len(avids), len(excluded_avids), len(checklist))

    checklist = [f"{avid}\n" for avid in checklist if avid not in historical_avids]

    print(len(avids), len(checklist))

    with open("checklist.csv", "w", encoding="utf-8-sig") as f:
        f.writelines(checklist)

    update_statuses(session, authenticity_token, excluded_avids)
    update_current(session, authenticity_token)
    download_current(WEEKS, session)


def update_statuses(session, authenticity_token, avids):
    """将需要排除的视频状态设为 0。"""
    data = {
        "utf8": "✓",
        "_method": "patch",
        "authenticity_token": authenticity_token,
        "status": "0",
        "desc": "",
    }

    for avid in avids:
        headers = {
            "Origin": BASE_URL,
            "Referer": f"{BASE_URL}/mad_result_tmps/{avid}/edit",
        }

        response = session.post(
            f"{BASE_URL}/mad_result_tmps/{avid}",
            headers=headers,
            data=data,
        )
        response.raise_for_status()
        print(avid, response.status_code)


def update_current(session, authenticity_token):
    """沿用列表页会话和 CSRF token，更新当前期。"""
    data = {
        "_method": "post",
        "authenticity_token": authenticity_token,
    }
    response = session.post(
        f"{BASE_URL}/mad_currents",
        headers={
            "Origin": BASE_URL,
            "Referer": f"{BASE_URL}/mad_result_tmps",
        },
        data=data,
    )
    response.raise_for_status()
    print("更新当前期", response.status_code)


def download_current(weeks, session=None):
    """下载指定期数的榜单，按既有十列格式保存为 Excel。"""
    if session is None:
        with create_session() as new_session:
            return download_current(weeks, new_session)

    response = session.get(
        f"{BASE_URL}/mad_ranking_results/{weeks}", timeout=30
    )
    response.raise_for_status()
    html = etree.HTML(response.content, parser=etree.HTMLParser(encoding="utf-8"))
    columns = ["rank", "av", "分", "title", "up", "点", "评", "弹", "藏", "币"]
    tables = [] if html is None else html.xpath("//table")
    table = next(
        (
            table
            for table in tables
            if ["".join(cell.itertext()).strip() for cell in table.xpath("./thead/tr/th")]
            == columns
        ),
        None,
    )
    if table is None:
        raise RuntimeError(f"第 {weeks} 期页面未找到榜单表格，请检查登录会话和期数")

    rows = []
    for row in table.xpath("./tbody/tr"):
        values = ["".join(cell.itertext()).strip() for cell in row.xpath("./td")]
        if len(values) != len(columns):
            raise ValueError(f"第 {weeks} 期榜单第 {len(rows) + 1} 行列数不正确")
        rows.append([
            value if index in (3, 4) else int(value)
            for index, value in enumerate(values)
        ])
    if not rows:
        raise RuntimeError(f"第 {weeks} 期榜单为空")

    workbook = Workbook()
    sheet = workbook.active
    if sheet is None:
        raise RuntimeError("无法创建 Excel 工作表")
    sheet.title = "Sheet1"
    sheet.append(columns)
    for values in rows:
        sheet.append(values)
        # 标题和 UP 主始终按文本写入，包括以等号开头的名称。
        for column in (4, 5):
            sheet.cell(sheet.max_row, column).data_type = "s"

    output = f"DATA/{weeks}期数据.xlsx"
    workbook.save(output)
    workbook.close()
    print(f"已保存 {output}，共 {len(rows)} 条记录")
    return output


if __name__ == "__main__":
    main()
